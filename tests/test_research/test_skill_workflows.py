"""Human-audited skill fixtures: provenance, edge cases and deterministic exports."""

import copy
import json
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfWriter
from pydantic import ValidationError

from openharness.utils.research_documents import parse_document
from openharness.utils.research_workflows.events import normalize_monitor
from openharness.utils.research_workflows.export import export_result
from openharness.utils.research_workflows.financial import calculate_financial
from openharness.utils.research_workflows.models import RESULT_TYPES
from openharness.utils.research_workflows.reports import calculate_deep, normalize_digest
from openharness.utils.session_files import SessionFiles

FIXTURES = Path(__file__).parents[1] / "fixtures" / "research_skills"
FUNCTIONS = {
    "financial": calculate_financial,
    "monitor": normalize_monitor,
    "digest": normalize_digest,
    "deep": calculate_deep,
}


def fixture(kind):
    return json.loads((FIXTURES / f"{kind}.json").read_text())


def compute(kind, data=None):
    return FUNCTIONS[kind](RESULT_TYPES[kind].model_validate(data or fixture(kind)))


def test_text_pdf_and_cross_page_positions():
    parsed = parse_document(FIXTURES / "annual-report.pdf")
    assert parsed["status"] == "ready" and len(parsed["pages"]) == 2
    assert "Consolidated year 2025" in parsed["pages"][0]["text"]
    assert "net_profit: 21" in parsed["pages"][1]["text"]
    assert "cash: 30" in parsed["pages"][0]["text"]
    assert len(parsed["document_hash"]) == 64
    txt = parse_document(FIXTURES / "annual-report.txt")
    assert txt["pages"][0]["page"] is None and txt["pages"][0]["start_line"] == 1


def test_scan_corrupt_encrypted_and_page_limit(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    assert parse_document(path)["status"] == "unsupported"
    writer.encrypt("secret")
    writer.write(path)
    assert "加密" in parse_document(path)["gaps"][0]
    path.write_bytes(b"%PDF-1.7\nbroken")
    with pytest.raises(ValueError):
        parse_document(path)
    monkeypatch.setattr("openharness.utils.research_documents.MAX_DOCUMENT_PAGES", 1)
    with pytest.raises(ValueError, match="1000"):
        parse_document(FIXTURES / "annual-report.pdf")


def test_financial_audited_values_and_units():
    result = compute("financial")
    assert result.status == "complete", result.gaps
    ratios = {row["metric"]: row for row in result.ratios}
    for name, expected in {
        "gross_margin": ".4",
        "net_margin": ".21",
        "debt_ratio": ".4",
        "current_ratio": "2",
    }.items():
        assert ratios[name]["value"] == Decimal(expected)
    assert ratios["simple_roe"]["value"] == Decimal(20) / Decimal(95)
    assert not any(row["annualized"] for row in result.ratios)
    assert all(row["state"] == "pass" for row in result.checks)
    data = fixture("financial")
    for amount in data["periods"][0]["items"].values():
        amount["value"] = str(Decimal(amount["value"]) / 10000)
        amount["unit"] = "亿元"
    assert compute("financial", data).ratios == result.ratios


@pytest.mark.parametrize(
    "key,value", [("opening_parent_equity", None), ("parent_equity", "-100"), ("revenue", "0")]
)
def test_missing_zero_and_negative_equity_never_become_valid_ratios(key, value):
    data = fixture("financial")
    data["periods"][0]["items"][key].update(
        value=value, missing_reason="未披露" if value is None else ""
    )
    result = compute("financial", data)
    metric = "gross_margin" if key == "revenue" else "simple_roe"
    ratio = next(row for row in result.ratios if row["metric"] == metric)
    assert ratio["value"] is None and ratio["limitation"]
    assert result.status == "partial"


def test_reconciliation_keeps_values_and_scopes_separate():
    data = fixture("financial")
    current = data["periods"][0]
    current["items"]["total_assets"]["value"] = "220"
    parent = copy.deepcopy(current)
    parent.update(scope="parent", label="2024母公司", start="2024-01-01", end="2024-12-31")
    data["periods"].append(parent)
    result = compute("financial", data)
    assert any(row["state"] == "fail" for row in result.checks)
    assert result.periods[0].items["total_assets"].value == Decimal(220)
    assert not result.changes  # Parent cannot be compared to consolidated.
    parent["scope"] = "consolidated"
    current["items"]["net_profit"]["value"] = "-20"
    parent["items"]["net_profit"]["value"] = "-10"
    change = next(
        row for row in compute("financial", data).changes if row["metric"] == "net_profit"
    )
    assert change["yoy"] is None and change["direction"] == "亏损扩大"


def test_monitor_boundaries_undated_dedup_and_conflicts():
    data = fixture("monitor")
    event = data["events"][0]
    event["published_at"] = "2026-09-28T12:00:00+08:00"  # Inclusive seven-day boundary.
    reprint = copy.deepcopy(event)
    reprint["references"][0]["locator"] = "fixture:reprint"
    reprint["impact"]["value"] = -1
    unknown = copy.deepcopy(event)
    unknown.update(event_key="undated", published_at=None)
    old = copy.deepcopy(event)
    old.update(event_key="old", published_at="2026-09-28T11:59:59+08:00")
    data["events"].extend([reprint, unknown, old])
    result = compute("monitor", data)
    assert len(result.events) == len(result.undated_events) == len(result.outside_window) == 1
    assert len(result.events[0].references) == 2 and result.events[0].impact.value is None
    assert "冲突" in result.events[0].impact.reason
    assert result.status == "partial" and str(result.as_of.tzinfo) == "Asia/Shanghai"


def test_monitor_no_events_is_distinct_from_failed_and_summary():
    data = fixture("monitor")
    data["events"] = []
    assert compute("monitor", data).status == "complete"
    data["channels"][0]["state"] = "failed"
    assert compute("monitor", data).status == "blocked"
    data["channels"][0]["state"] = "summary_only"
    assert compute("monitor", data).status == "partial"


def test_digest_comparability_and_unknown_prior_rating():
    data = fixture("digest")
    report = data["reports"][0]
    report["rating_change"] = "上调"
    other = copy.deepcopy(report)
    other.update(institution="样例证券B", title="另一机构研报")
    other["predictions"][0]["amount"].update(value="20000", unit="万元")
    different = copy.deepcopy(other)
    different["predictions"][0]["basis"] = "母公司"
    data["reports"].extend([other, different])
    result = compute("digest", data)
    assert result.reports[0].rating_change is None
    assert len(result.comparison) == 2
    assert len(result.comparison[0]["predictions"]) == 2
    assert {row["value_yuan"] for row in result.comparison[0]["predictions"]} == {
        Decimal(200000000)
    }
    report["summary_only"] = True
    assert compute("digest", data).status == "partial"


def test_deep_three_scenarios_and_reproducible_sensitivity():
    result = compute("deep")
    assert result.status == "complete", result.gaps
    assert len(result.forecasts) == 9 and len(result.sensitivity) == 27
    first = result.forecasts[0]["values"]
    assert first["revenue"] == Decimal(1100000)
    assert first["operating_profit"] == Decimal(285000)
    assert first["net_profit"] == Decimal(228750)
    assert first["eps"] == Decimal(".22875")
    assert result.sensitivity[0]["parent_net_profit"] == Decimal(208800)
    assert result.sensitivity[8]["parent_net_profit"] == Decimal(249300)
    assert result.sensitivity[4]["parent_net_profit"] == first["parent_net_profit"]
    assert compute("deep").model_dump() == result.model_dump()
    varied = fixture("deep")
    for scenario in varied["scenarios"]:
        if scenario["name"] == "base":
            continue
        for year in scenario["years"]:
            year["assumptions"]["growth"]["value"] = (
                ".2" if scenario["name"] == "optimistic" else "-.1"
            )
            year["assumptions"]["gross_margin"]["value"] = (
                ".45" if scenario["name"] == "optimistic" else ".35"
            )
    varied = compute("deep", varied)
    profits = {
        row["scenario"]: row["values"]["net_profit"]
        for row in varied.forecasts
        if row["year"] == 2026
    }
    assert profits == {
        "base": Decimal(228750),
        "optimistic": Decimal(292500),
        "cautious": Decimal(157500),
    }
    data = fixture("deep")
    data["scenarios"][0]["years"][0]["assumptions"]["growth"]["value"] = None
    partial = compute("deep", data)
    assert partial.status == "partial" and not any(
        row["scenario"] == "base" for row in partial.forecasts
    )
    data = fixture("deep")
    data["base_year"] = 2024
    for scenario in data["scenarios"]:
        for index, year in enumerate(scenario["years"]):
            year["year"] = 2025 + index
    assert not compute("deep", data).forecasts


def test_unknown_values_and_sector_fail_fast():
    data = fixture("financial")
    data["company"]["financial_sector"] = True
    with pytest.raises(ValidationError, match="非金融"):
        compute("financial", data)
    data = fixture("financial")
    data["periods"][0]["items"]["cash"]["references"] = []
    with pytest.raises(ValidationError, match="定位"):
        compute("financial", data)


@pytest.mark.parametrize("kind", RESULT_TYPES)
def test_exports_values_sources_and_session_artifacts(kind, tmp_path):
    result = compute(kind)
    session = tmp_path / "session"
    exported = export_result(result, tmp_path / "outputs", session, "task1")
    assert {Path(path).suffix for path in exported["files"]} == {".md", ".json", ".docx", ".xlsx"}
    data = json.loads((tmp_path / "outputs" / f"{kind}.json").read_text())
    markdown = (tmp_path / "outputs" / f"{kind}.md").read_text()
    assert data == result.model_dump(mode="json")
    assert "fixture:annual-report.pdf" in markdown and "[R1]" in markdown
    doc = Document(tmp_path / "outputs" / f"{kind}.docx")
    assert any("资料说明" in row.text for row in doc.paragraphs)
    if kind in {"financial", "deep"}:
        assert doc.tables
    wb = load_workbook(tmp_path / "outputs" / f"{kind}.xlsx", data_only=True)
    if kind == "financial":
        row = next(row for row in wb["ratios"].values if row[0] == "[0].value")
        assert row[1] == 0.4  # Precalculated numeric cells, independent of Excel recalculation.
    if kind == "deep":
        row = next(row for row in wb["forecasts"].values if row[0] == "[0].values.revenue")
        assert row[1] == 1100000
        assert "formulas" in str(list(wb["forecasts"].values))
    files = SessionFiles(session)
    assert len(files.list("artifacts")) == 4
    for artifact in exported["artifacts"]:
        meta, path = files.artifact(artifact["id"])
        assert path.is_file() and meta["task_id"] == "task1"
    with pytest.raises(FileNotFoundError):
        SessionFiles(tmp_path / "other-session").artifact(exported["artifacts"][0]["id"])


def test_cross_session_source_ids_rejected_before_export(tmp_path):
    data = fixture("digest")
    data["inputs"][0]["source_id"] = "previous-session-source"
    with pytest.raises(ValueError, match="重新登记"):
        export_result(compute("digest", data), tmp_path / "out", tmp_path / "session")
    assert not (tmp_path / "out").exists()


def test_attachment_storage_opaque_paths_and_restart(tmp_path):
    storage = SessionFiles(tmp_path)
    upload = storage.upload("../annual.pdf", (FIXTURES / "annual-report.pdf").read_bytes())
    assert upload["name"] == "annual.pdf" and upload["status"] == "ready"
    restarted = SessionFiles(tmp_path)
    assert restarted.list("attachments") == [upload]
    assert "parsed.json" in restarted.describe([upload["id"]])
    with pytest.raises(ValueError):
        storage.attachment("../../etc/passwd")
    with pytest.raises(ValueError, match="仅支持"):
        storage.upload("image.png", b"not supported")
    storage.delete_attachment(upload["id"])
    assert not restarted.list("attachments")


def test_restated_comparative_takes_priority_and_missing_currency_stays_unknown():
    data = fixture("financial")
    prior = copy.deepcopy(data["periods"][0])
    prior.update(label="2024原数", start="2024-01-01", end="2024-12-31")
    restated = copy.deepcopy(prior)
    restated.update(label="2024重述", restated=True)
    restated["items"]["revenue"]["value"] = "80"
    data["periods"].extend([prior, restated])
    result = compute("financial", data)
    revenue_changes = [row for row in result.changes if row["metric"] == "revenue"]
    assert len(revenue_changes) == 1 and revenue_changes[0]["prior"] == "2024重述"
    assert revenue_changes[0]["yoy"] == Decimal(".25")
    data = fixture("financial")
    data["periods"][0].pop("currency")
    result = compute("financial", data)
    assert result.periods[0].currency is None and result.status == "partial"
    assert any("币种未知" in gap for gap in result.gaps)


def test_no_reliable_share_count_means_no_eps():
    data = fixture("deep")
    for scenario in data["scenarios"]:
        for year in scenario["years"]:
            year["assumptions"]["shares"].update(origin="analyst", references=[])
    result = compute("deep", data)
    assert all(row["values"]["eps"] is None for row in result.forecasts)
    assert result.status == "partial"


def test_plugin_script_limits_workflow_to_its_own_kind(capsys):
    from openharness.utils.research_workflows.cli import main

    with pytest.raises(SystemExit) as exc:
        main(["schema", "deep"], skill_kind="financial")
    assert exc.value.code == 2 and "invalid choice" in capsys.readouterr().err
    main(["schema", "financial"], skill_kind="financial")
    assert json.loads(capsys.readouterr().out)["title"] == "FinancialResult"
