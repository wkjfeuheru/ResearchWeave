"""Read deliverables for validation and judging, including actual Office contents."""

import hashlib
import json
from pathlib import Path


def read_deliverable(path: str | Path) -> tuple[str, dict[str, object]]:
    path = Path(path)
    if path.stat().st_size > 2_000_000:
        raise ValueError("产物超过评分读取上限")
    data = path.read_bytes()
    if path.suffix in {".md", ".json"}:
        text = data.decode("utf-8")
        if path.suffix == ".json":
            json.loads(text)
    elif path.suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        )
    elif path.suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=False)
        try:
            text = "\n".join(
                f"{sheet.title}: " + " | ".join(str(v) if v is not None else "" for v in row)
                for sheet in book
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            book.close()
    else:
        raise ValueError("不支持的交付格式")
    return text, {
        "valid": bool(text.strip()),
        "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "format": path.suffix,
    }
