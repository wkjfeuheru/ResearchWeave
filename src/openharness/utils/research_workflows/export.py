"""Readable, consistent research exports from validated structured results."""

from __future__ import annotations

import json
import re
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from openharness.utils.fs import atomic_write_text
from openharness.utils.session_files import SessionFiles
from .financial import CORE_ITEMS

KIND_LABELS = {
    "financial": "财报穿透解析",
    "monitor": "舆情与公告监控",
    "digest": "研报精读与摘要",
    "deep": "深度投研报告",
}
METRIC_LABELS = {
    "gross_margin": "毛利率",
    "net_margin": "合并净利率",
    "debt_ratio": "资产负债率",
    "current_ratio": "流动比率",
    "cash_to_profit": "经营现金流/合并净利润",
    "simple_roe": "简化ROE（非加权平均）",
}


def collect_references(value):
    if isinstance(value, dict):
        if "locator" in value and "title" in value:
            yield value
        for item in value.values():
            yield from collect_references(item)
    elif isinstance(value, list):
        for item in value:
            yield from collect_references(item)


def display(value):
    if value is None:
        return "未知/不可计算"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def render_markdown(data: dict, session_directory: Path | None = None) -> str:
    references = []
    for ref in collect_references(data):
        if ref not in references:
            references.append(ref)
    memory = {}
    if session_directory is not None and (session_directory / "state.json").is_file():
        memory = json.loads((session_directory / "state.json").read_text())
    for ref in references:
        for key, pool in (("source_id", "sources"), ("evidence_id", "evidence_pool")):
            if ref.get(key) and ref[key] not in memory.get(pool, {}):
                raise ValueError(f"{ref[key]}不属于当前会话；导入资料须重新登记来源/证据")
        if ref.get("source_id") and ref.get("evidence_id"):
            if memory["evidence_pool"][ref["evidence_id"]]["source_id"] != ref["source_id"]:
                raise ValueError("来源与证据不对应，须核对引用")

    def refs(items):
        return " ".join(f"[R{references.index(ref) + 1}]" for ref in items)

    def claim(item):
        return item["text"] + " " + refs(item["references"])

    lines = [
        f"# {KIND_LABELS[data['kind']]}：{data['company']['name']}（{data['company']['code']}）",
        f"状态：{data['status']}；资料截止：{data['as_of']}",
        "",
    ]
    kind = data["kind"]
    if kind == "financial":
        for period in data["periods"]:
            lines.extend(
                [
                    f"## {period['label']} / {period['scope']} / {period['currency']}",
                    f"期间：{period['start']} 至 {period['end']}；重述：{period['restated']}",
                    "| 科目 | 原披露数值 | 单位 | 依据/缺口 |",
                    "|---|---:|---|---|",
                ]
            )
            for key, item in period["items"].items():
                lines.append(
                    f"| {CORE_ITEMS[key]} | {display(item['value'])} | {item['unit']} | {refs(item['references'])} {item['missing_reason']} |"
                )
        lines.extend(
            [
                "## 财务比率",
                "| 期间/口径 | 指标 | 比率值（未年化） | 公式 | 限制 |",
                "|---|---|---:|---|---|",
            ]
        )
        for row in data["ratios"]:
            lines.append(
                f"| {row['period']}/{row['scope']} | {METRIC_LABELS[row['metric']]} | {display(row['value'])} | {row['formula']} | {row['limitation']} {refs(row['references'])} |"
            )
        lines.append("## 勾稽校验")
        for row in data["checks"]:
            lines.append(
                f"- {row['period']}/{row['scope']} {row['check']}: {row['state']}；差额 {display(row['residual_yuan'])} 元，容差 {row['tolerance_yuan']} 元。{refs(row['references'])}"
            )
        lines.append("## 同比与亏损方向")
        for row in data["changes"]:
            lines.append(
                f"- {row['period']}/{row['scope']} {CORE_ITEMS[row['metric']]}：{row['direction']}，差额 {row['difference_yuan']} 元；同比 {display(row['yoy'])}。{row['limitation']} {refs(row['references'])}"
            )
        lines.append("## 关键附注核查")
        for finding in data["notes"]:
            lines.extend(
                [
                    f"### {finding['topic']}",
                    f"事实：{finding['fact']} {refs(finding['references'])}",
                    f"解释：{finding['interpretation']}",
                ]
            )
    elif kind == "monitor":
        lines.append(f"检索窗口：{data['window_start']} 至 {data['as_of']}；覆盖仅限实际取得资料。")
        for key, heading in (
            ("events", "窗口内事件"),
            ("undated_events", "发布日期未知"),
            ("outside_window", "窗口外资料"),
        ):
            lines.append(f"## {heading}")
            if not data[key]:
                lines.append(
                    "未取得此类事件；不代表不存在事件。"
                    if data["status"] != "complete"
                    else "本次查询未发现此类事件。"
                )
            for event in data[key]:
                lines.extend(
                    [
                        f"### {event['title']} / {event['category']}",
                        f"发布：{event['published_at']}；发生：{event['occurred_at']}；仅摘要：{event['summary_only']}",
                        event["summary"] + " " + refs(event["references"]),
                    ]
                )
                for score_key, label in (("sentiment", "文本情绪"), ("impact", "基本面影响")):
                    score = event[score_key]
                    lines.append(
                        f"- {label}：{display(score['value'])}（置信度 {score['confidence']}）；{score['reason']}"
                    )
        lines.append("## 渠道状态")
        lines.extend(f"- {row['name']}: {row['state']} {row['detail']}" for row in data["channels"])
    elif kind == "digest":
        for report in data["reports"]:
            lines.extend(
                [
                    f"## {report['title']}",
                    f"机构：{report['institution']}；作者：{'、'.join(report['authors'])}；日期：{report['published_at']}；仅摘要：{report['summary_only']}",
                    refs(report["references"]),
                ]
            )
            for key, heading in (
                ("views", "核心观点"),
                ("arguments", "核心论据"),
                ("assumptions", "预测假设"),
                ("risks", "风险提示"),
            ):
                lines.append(f"### {heading}")
                lines.extend("- " + claim(item) for item in report[key])
            lines.append("### 原研报盈利预测（非实际业绩）")
            for prediction in report["predictions"]:
                amount = prediction["amount"]
                lines.append(
                    f"- {prediction['year']} {prediction['metric']}: {display(amount['value'])} {amount['unit']} {prediction['currency']}；{prediction['basis']} {refs(amount['references'])}"
                )
            rating = report["rating"]
            lines.append(
                f"评级：{rating['current']}；前次：{rating['previous']}；变动：{report['rating_change'] or '未知'} {refs(rating['references'])}"
            )
            target = report["target_price"]
            lines.append(
                f"原研报目标价：{display(target['value'])} {report['target_currency']} {refs(target['references'])}"
                if target
                else "原研报目标价：未知"
            )
        lines.extend(
            ["## 可比预测对照", "以下仅合并同年度、币种、指标和口径，不平均不同机构评级。"]
        )
        for group in data["comparison"]:
            lines.append(
                f"- {group['year']} / {group['metric']} / {group['currency']} / {group['basis']}: "
                + "; ".join(
                    f"{item['institution']} ({item['published_at']}) {display(item['value_yuan'])} 元"
                    for item in group["predictions"]
                )
            )
    else:
        for section in data["sections"]:
            lines.append(f"## {section['title']}")
            lines.extend(claim(item) for item in section["paragraphs"])
            lines.extend("缺口：" + item for item in section["gaps"])
        lines.extend(
            ["## 三年情景盈利预测", "预测基于显式假设，非已实现业绩；未生成自主评级或目标价。"]
        )
        baseline = data["base_revenue"]
        lines.append(
            f"基期：{data['base_year']}年度合并收入 {display(baseline['value'])} {baseline['unit']} {data['base_currency']}；{refs(baseline['references'])}"
        )
        lines.append(
            "| 情景 | 年度 | 收入（元） | 营业利润（元） | 税前利润（元） | 合并净利润（元） | 归母净利润（元） | EPS |"
        )
        lines.append("|---|---|---:|---:|---:|---:|---:|---:|")
        for row in data["forecasts"]:
            values = row["values"]
            lines.append(
                f"| {row['scenario']} | {row['year']} | {values['revenue']} | {values['operating_profit']} | {values['pretax_profit']} | {values['net_profit']} | {values['parent_net_profit']} | {display(values['eps'])} |"
            )
        lines.append("### 预测假设与依据")
        for scenario in data["scenarios"]:
            for year in scenario["years"]:
                for key, assumption in year["assumptions"].items():
                    lines.append(
                        f"- {scenario['name']}/{year['year']} {key}: {display(assumption['value'])} ({assumption['origin']})；{assumption['rationale']} {refs(assumption['references'])}"
                    )
        lines.append("### 单年度敏感性（其他条件保持基准）")
        lines.extend(
            f"- {row['year']}: 增速变动 {row['growth_shift']}，毛利率变动 {row['gross_margin_shift']}，归母净利润 {row['parent_net_profit']} 元"
            for row in data["sensitivity"]
        )
    lines.append("## 缺口与限制")
    lines.extend("- " + gap for gap in data["gaps"])
    if not data["gaps"]:
        lines.append("本次结构化检查未登记额外缺口；不代表所有事实已核验。")
    lines.append("## 资料说明")
    for index, ref in enumerate(references, 1):
        lines.append(
            f"- [R{index}] {ref['title']}；{ref['locator']}；页码 {ref.get('page') or '不适用/未知'}；发布日期 {ref.get('published_at') or '未知'}"
        )
    # GFM table rows must remain adjacent; prose blocks get blank-line separation.
    text = "\n".join(line if line.startswith("|") else "\n" + line + "\n" for line in lines)

    def replace_evidence(match):
        evidence = memory.get("evidence_pool", {}).get(match[1])
        source = memory.get("sources", {}).get(evidence.get("source_id")) if evidence else None
        if source is None:
            raise ValueError("报告包含无法解析的证据编号，须重新登记来源")
        return f"（{source['title']}，{source['locator']}）"

    return re.sub(r"\[E:([^\]]+)\]", replace_evidence, text)


def flatten(value, prefix=""):
    if isinstance(value, dict):
        for key, item in value.items():
            yield from flatten(item, f"{prefix}.{key}" if prefix else key)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from flatten(item, f"{prefix}[{index}]")
    else:
        yield prefix, value


def export_result(
    result, directory: Path, session_directory: Path | None = None, task_id: str | None = None
) -> dict:
    data = result.model_dump(mode="json")
    raw_data = result.model_dump(mode="python")
    markdown = render_markdown(data, session_directory)
    directory = directory.resolve()
    directory.mkdir(parents=True, exist_ok=True)
    prefix = data["kind"]
    paths = [directory / (prefix + suffix) for suffix in (".md", ".json", ".docx", ".xlsx")]
    atomic_write_text(paths[0], markdown)
    atomic_write_text(paths[1], json.dumps(data, ensure_ascii=False, indent=2))
    document = Document()
    table = None
    for line in markdown.splitlines():
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                continue
            if table is None:
                table = document.add_table(rows=0, cols=len(cells))
                table.style = "Table Grid"
            row = table.add_row().cells
            for cell, value in zip(row, cells):
                cell.text = value
            continue
        if line.strip():
            table = None
        if line.startswith("#"):
            size = len(line) - len(line.lstrip("#"))
            document.add_heading(line.lstrip("# "), level=min(size, 4))
        elif line.strip():
            document.add_paragraph(line)
    document.save(paths[2])
    workbook = Workbook()
    workbook.remove(workbook.active)
    for key in (
        "company",
        "periods",
        "ratios",
        "checks",
        "changes",
        "notes",
        "channels",
        "events",
        "undated_events",
        "outside_window",
        "reports",
        "comparison",
        "base_revenue",
        "scenarios",
        "forecasts",
        "sensitivity",
        "gaps",
        "inputs",
    ):
        if key not in data:
            continue
        sheet = workbook.create_sheet(key[:31])
        sheet.append(["字段/位置", "数值或说明", "精确值（Decimal；避免Excel精度截断）"])
        for name, value in flatten(raw_data[key]):
            exact = str(value) if isinstance(value, Decimal) else None
            if isinstance(value, Decimal):
                value = float(value)
            if not isinstance(value, (str, int, float, bool, type(None))):
                value = value.isoformat() if hasattr(value, "isoformat") else str(value)
            if isinstance(value, str):
                # Preserve literal external text, not spreadsheet formulas.
                sheet.append([name, value, exact])
                sheet.cell(sheet.max_row, 2).data_type = "s"
            else:
                sheet.append([name, value, exact])
            if exact is not None:
                sheet.cell(sheet.max_row, 3).data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 55
        sheet.column_dimensions["B"].width = 80
        sheet.column_dimensions["C"].width = 38
    workbook.save(paths[3])
    artifacts = []
    if session_directory is not None:
        storage = SessionFiles(session_directory)
        artifacts = [
            storage.register(path, task_id=task_id, status=data["status"], kind=data["kind"])
            for path in paths
        ]
    return {
        "status": data["status"],
        "gaps": data["gaps"],
        "files": [str(path) for path in paths],
        "artifacts": artifacts,
    }
