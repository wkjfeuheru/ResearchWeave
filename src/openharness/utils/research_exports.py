"""Generic cited-document and spreadsheet writers; no skill business dispatch."""

from __future__ import annotations
from pydantic import BaseModel
from typing import Any, Iterator, Callable, Iterable

import json
import re
import os
from decimal import Decimal
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from openharness.utils.fs import atomic_write_text
from openharness.utils.session_files import SessionFiles


def collect_references(value: object) -> Iterator[dict[str, Any]]:
    if isinstance(value, dict):
        if "locator" in value and "title" in value:
            yield value
        for item in value.values():
            yield from collect_references(item)
    elif isinstance(value, list):
        for item in value:
            yield from collect_references(item)


def display(value: object) -> str:
    if value is None:
        return "未知/不可计算"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


class ReportContext:
    """Shared reference checks and source navigation for plugin-owned report bodies."""

    def __init__(self, data: dict[str, Any], session_directory: Path | None = None) -> None:
        self.data = data
        self.references: list[dict[str, Any]] = []
        for ref in collect_references(data):
            if ref not in self.references:
                self.references.append(ref)
        self.memory: dict[str, Any] = {}
        if session_directory is not None and (session_directory / "state.json").is_file():
            self.memory = json.loads((session_directory / "state.json").read_text())
        for ref in self.references:
            if os.environ.get("OPENHARNESS_ISOLATED_EXPORT"):
                continue  # Candidate IDs are verified by the main agent on the host.
            for key, pool in (("source_id", "sources"), ("evidence_id", "evidence_pool")):
                if ref.get(key) and ref[key] not in self.memory.get(pool, {}):
                    raise ValueError(f"{ref[key]}不属于当前会话；导入资料须重新登记来源/证据")
            if ref.get("source_id") and ref.get("evidence_id"):
                if (
                    self.memory["evidence_pool"][ref["evidence_id"]]["source_id"]
                    != ref["source_id"]
                ):
                    raise ValueError("来源与证据不对应，须核对引用")

    def refs(self, items: list[dict[str, Any]]) -> str:
        return " ".join(f"[R{self.references.index(ref) + 1}]" for ref in items)

    def claim(self, item: dict[str, Any]) -> str:
        return str(item["text"]) + " " + self.refs(item["references"])

    def header(self, title: str) -> list[str]:
        data = self.data
        return [
            f"# {title}：{data['company']['name']}（{data['company']['code']}）",
            f"状态：{data['status']}；资料截止：{data['as_of']}",
            "",
        ]

    def finish(self, lines: list[str]) -> str:
        lines.append("## 缺口与限制")
        lines.extend("- " + gap for gap in self.data["gaps"])
        if not self.data["gaps"]:
            lines.append("本次结构化检查未登记额外缺口；不代表所有事实已核验。")
        lines.append("## 资料说明")
        for index, ref in enumerate(self.references, 1):
            lines.append(
                f"- [R{index}] {ref['title']}；{ref['locator']}；页码 {ref.get('page') or '不适用/未知'}；发布日期 {ref.get('published_at') or '未知'}"
            )
        # GFM table rows must remain adjacent; prose blocks get blank-line separation.
        text = "\n".join(line if line.startswith("|") else "\n" + line + "\n" for line in lines)

        def replace_evidence(match: re.Match[str]) -> str:
            evidence = self.memory.get("evidence_pool", {}).get(match[1])
            source = (
                self.memory.get("sources", {}).get(evidence.get("source_id")) if evidence else None
            )
            if source is None:
                raise ValueError("报告包含无法解析的证据编号，须重新登记来源")
            return f"（{source['title']}，{source['locator']}）"

        return re.sub(r"\[E:([^\]]+)\]", replace_evidence, text)


def flatten(value: object, prefix: str = "") -> Iterator[tuple[str, object]]:
    if isinstance(value, dict):
        for key, item in value.items():
            yield from flatten(item, f"{prefix}.{key}" if prefix else key)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from flatten(item, f"{prefix}[{index}]")
    else:
        yield prefix, value


def export_result(
    result: BaseModel,
    directory: Path,
    session_directory: Path | None = None,
    task_id: str | None = None,
    *,
    render_markdown: Callable[[dict[str, Any], Path | None], str],
    sheets: Iterable[str],
) -> dict[str, object]:
    data = result.model_dump(mode="json")
    raw_data = result.model_dump(mode="python")
    markdown = render_markdown(data, session_directory)
    directory = directory.resolve()
    directory.mkdir(parents=True, exist_ok=True)
    prefix = data["kind"]
    paths = [directory / (prefix + suffix) for suffix in (".md", ".json", ".docx", ".xlsx")]
    atomic_write_text(paths[0], markdown)
    atomic_write_text(paths[1], json.dumps(data, ensure_ascii=False, indent=2))
    document = Document()
    table = None
    for line in markdown.splitlines():
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                continue
            if table is None:
                table = document.add_table(rows=0, cols=len(cells))
                table.style = "Table Grid"
            row = table.add_row().cells
            for cell, cell_text in zip(row, cells):
                cell.text = cell_text
            continue
        if line.strip():
            table = None
        if line.startswith("#"):
            size = len(line) - len(line.lstrip("#"))
            document.add_heading(line.lstrip("# "), level=min(size, 4))
        elif line.strip():
            document.add_paragraph(line)
    document.save(paths[2])
    workbook = Workbook()
    active_sheet = workbook.active
    if active_sheet is not None:
        workbook.remove(active_sheet)
    for key in sheets:
        if key not in data:
            continue
        sheet = workbook.create_sheet(key[:31])
        sheet.append(["字段/位置", "数值或说明", "精确值（Decimal；避免Excel精度截断）"])
        for name, value in flatten(raw_data[key]):
            exact = str(value) if isinstance(value, Decimal) else None
            if isinstance(value, Decimal):
                value = float(value)
            if not isinstance(value, (str, int, float, bool, type(None))):
                value = value.isoformat() if hasattr(value, "isoformat") else str(value)
            if isinstance(value, str):
                # Preserve literal external text, not spreadsheet formulas.
                sheet.append([name, value, exact])
                sheet.cell(sheet.max_row, 2).data_type = "s"
            else:
                sheet.append([name, value, exact])
            if exact is not None:
                sheet.cell(sheet.max_row, 3).data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 55
        sheet.column_dimensions["B"].width = 80
        sheet.column_dimensions["C"].width = 38
    workbook.save(paths[3])
    artifacts = []
    import os

    if session_directory is not None and not os.environ.get("OPENHARNESS_ISOLATED_EXPORT"):
        execution = json.loads(os.environ.get("OPENHARNESS_RESEARCH_EXECUTION", "null"))
        if execution and (
            execution["task_id"] != task_id
            or Path(os.environ.get("OPENHARNESS_RESEARCH_SESSION_DIR", "")).resolve()
            != session_directory.resolve()
        ):
            raise ValueError("Export execution belongs to another task or session")
        storage = SessionFiles(session_directory)
        artifacts = [
            storage.register(
                path, task_id=task_id, status=data["status"], kind=data["kind"], execution=execution
            )
            for path in paths
        ]
    return {
        "status": data["status"],
        "gaps": data["gaps"],
        "files": [str(path) for path in paths],
        "artifacts": artifacts,
    }
